"""Extract Zensus 2022 employed persons by age class and sex per Kreis from the Regionaltabelle.

ADR-0137. The popsim employment grid splits each Kreis's employed persons (census level
``ERWERBSTAT_KURZ_STP__11_{M,W}``, cleancensus ``kreis_erwerbsstatus.parquet``) across five age
groups. Its previous shape reference (``zensus2022_employment_by_age_ref.csv``) was exact for the
three kreisfreie Staedte only, so the five Landkreise used a national shape. The Zensus 2022
Regionaltabelle "Bildung und Erwerbstaetigkeit" publishes the same quantity -- employed persons
("Erwerbstaetige", sample-based STP results) by age class -- for EVERY Kreis, by sex, at the same
reference date (2022-05-15) as the census level. This script writes that table for the requested
regions as the local-only reference ``zensus2022_employed_by_age_kreis.csv`` (not committed, per the
``.gitignore`` data policy; Data Registry record ``zensus2022_employed_by_age_kreis``).

Source: ``Regionaltabelle_Bildung_Erwerbstaetigkeit.xlsx`` (Statistische Aemter des Bundes und der
Laender, "Zensus 2022 - Ausgewaehlte Zensusergebnisse", Datenlizenz Deutschland - Namensnennung -
Version 2.0), local copy under ``<cleancensus>/data/raw/regionaltabellen/``:

* sheet ``CSV-ET_Alter``: one row per region (``_RS`` = ARS: ``00`` Germany, 2-digit Land,
  5-digit Kreis, 12-digit Gemeinde) with the counts in ``ET_ALTER_ERW__<i>_<M|W>`` for the
  seven age classes ``i = 1..7``;
* sheet ``ET Alter``: the same table for reading, whose class header names what ``i`` means.
  It is checked against :data:`AGE_CLASSES` so that a reordered release fails instead of
  silently mislabelling every class.

The Zensus withholds small cells ("/"); a withheld or non-numeric count for a REQUESTED region
is an error, since a shape with a missing class would be wrong, not approximate.

Output CSV (``#``-prefixed provenance header, then): ``region,region_name,age_class,sex,
employed_persons`` with ``sex`` in ``M``/``F`` (the source's ``W``) and one row per region, class
and sex.

Usage::

    python scripts/extract_zensus2022_employed_by_age.py \\
        --regionaltabelle <cleancensus>/data/raw/regionaltabellen/Regionaltabelle_Bildung_Erwerbstaetigkeit.xlsx \\
        --out eqasim-data/data/braunschweig/popsim/zensus2022_employed_by_age_kreis.csv
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
from pathlib import Path
from typing import Iterable, List

import openpyxl
import pandas as pd

# Run as ``python scripts/<this file>``, the repository root is not on sys.path; the comparison
# mode imports braunschweig.popsim (same pattern as scripts/extract_commute_day_diagnostics.py).
REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

#: The seven age classes of ``ET Alter`` / ``ET_ALTER_ERW__1..7``, in source order.
AGE_CLASSES = ("15-19", "20-29", "30-39", "40-49", "50-59", "60-67", "68+")

#: The eight ZGB Kreise and Germany (``00``), the national row a missing Kreis falls back to.
DEFAULT_REGIONS = ("03101", "03102", "03103", "03151", "03153", "03154", "03157", "03158", "00")

MACHINE_SHEET = "CSV-ET_Alter"
READABLE_SHEET = "ET Alter"
REFERENCE_DATE = "2022-05-15"
_SOURCE_SEX = (("M", "M"), ("W", "F"))
_OUTPUT_COLUMNS = ["region", "region_name", "age_class", "sex", "employed_persons"]


def _normalised_class_label(label) -> str:
    """``"15 - 19 "`` -> ``"15-19"``, ``"68 und aelter"`` -> ``"68+"``."""
    text = re.sub(r"\s+", "", str(label))
    return re.sub(r"und(ä|ae)lter$", "+", text)


def _readable_class_labels(workbook) -> List[str]:
    """The class labels of the readable sheet, in order (the row that starts with ``15``)."""
    for row in workbook[READABLE_SHEET].iter_rows(min_row=1, max_row=12, values_only=True):
        labels = [_normalised_class_label(value) for value in row if value is not None]
        if labels and labels[0] == AGE_CLASSES[0]:
            return labels
    raise ValueError(f"sheet {READABLE_SHEET!r} has no age-class header row starting with "
                     f"{AGE_CLASSES[0]!r}; the Regionaltabelle layout changed.")


def extract_employed_by_age(workbook_path, regions: Iterable[str]) -> pd.DataFrame:
    """Employed persons by age class and sex for ``regions`` (ARS codes), one row per triple.

    Raises ``ValueError`` when the class order differs from :data:`AGE_CLASSES`, when a
    requested region is absent or duplicated, or when one of its counts is withheld or
    non-numeric.
    """
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    labels = _readable_class_labels(workbook)
    if tuple(labels) != AGE_CLASSES:
        raise ValueError(f"the age classes of sheet {READABLE_SHEET!r} are {labels}, not the "
                         f"expected {list(AGE_CLASSES)}; the ET_ALTER_ERW__<i> columns would be "
                         "mislabelled.")
    rows = workbook[MACHINE_SHEET].iter_rows(values_only=True)
    header = list(next(rows))
    position = {name: index for index, name in enumerate(header)}
    wanted = [str(region) for region in regions]
    found = {}
    for row in rows:
        region = str(row[position["_RS"]])
        if region in wanted:
            if region in found:
                raise ValueError(f"region {region} appears twice in sheet {MACHINE_SHEET!r}.")
            found[region] = row
    missing = [region for region in wanted if region not in found]
    if missing:
        raise ValueError(f"regions {missing} are not in sheet {MACHINE_SHEET!r} of {workbook_path}.")

    records = []
    for region in wanted:
        row = found[region]
        for class_index, age_class in enumerate(AGE_CLASSES, start=1):
            for source_sex, sex in _SOURCE_SEX:
                column = f"ET_ALTER_ERW__{class_index}_{source_sex}"
                value = row[position[column]]
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    raise ValueError(
                        f"region {region} ({row[position['Name']]}): {column} is {value!r}, not a "
                        "count -- the Zensus withholds such cells, and a shape without this "
                        "class would be wrong.")
                records.append((region, str(row[position["Name"]]), age_class, sex, int(value)))
    frame = pd.DataFrame(records, columns=_OUTPUT_COLUMNS)
    frame["employed_persons"] = frame["employed_persons"].astype("int64")
    return frame


def _sha256(path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def write_reference(frame: pd.DataFrame, destination, source_path) -> None:
    """Write ``frame`` with a ``#`` provenance header naming the source file and its digest."""
    os.makedirs(os.path.dirname(os.path.abspath(destination)), exist_ok=True)
    provenance = [
        "Zensus 2022 employed persons (Erwerbstaetige, sample-based STP results) by age class "
        "and sex per region.",
        f"source: {Path(source_path).name}, sheet {MACHINE_SHEET} (classes checked against "
        f"sheet {READABLE_SHEET}); reference date {REFERENCE_DATE}.",
        f"source sha256: {_sha256(source_path)}",
        "publisher: Statistische Aemter des Bundes und der Laender; licence: Datenlizenz "
        "Deutschland - Namensnennung - Version 2.0 (dl-de/by-2-0).",
        "written by scripts/extract_zensus2022_employed_by_age.py (ADR-0137).",
    ]
    with open(destination, "w", encoding="utf-8", newline="") as handle:
        for line in provenance:
            handle.write(f"# {line}\n")
        frame.to_csv(handle, index=False)


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--regionaltabelle", required=True,
                        help="path of Regionaltabelle_Bildung_Erwerbstaetigkeit.xlsx")
    parser.add_argument("--out", required=True, help="destination CSV")
    parser.add_argument("--regions", nargs="+", default=list(DEFAULT_REGIONS),
                        help="ARS codes to extract (default: the 8 ZGB Kreise and Germany '00')")
    parser.add_argument("--compare-with",
                        help="the previous reference zensus2022_employment_by_age_ref.csv; prints "
                             "(and with --comparison-out writes) the change per Kreis, sex and group")
    parser.add_argument("--comparison-out", help="destination CSV of the comparison")
    args = parser.parse_args(argv)
    frame = extract_employed_by_age(args.regionaltabelle, args.regions)
    write_reference(frame, args.out, args.regionaltabelle)
    print(f"[extract_zensus2022_employed_by_age] wrote {len(frame)} rows for "
          f"{frame['region'].nunique()} regions to {args.out}")
    if args.compare_with:
        # Imported here: the extraction itself needs nothing from the package.
        from braunschweig.popsim import zensus_employment_age

        kreise = [region for region in args.regions if region != zensus_employment_age.NATIONAL_REGION]
        comparison = zensus_employment_age.compare_age_shapes(args.compare_with, args.out, kreise)
        print(comparison.round(2).to_string(index=False))
        if args.comparison_out:
            comparison.to_csv(args.comparison_out, index=False)


if __name__ == "__main__":
    main()
